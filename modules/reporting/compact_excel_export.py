"""Safe streaming export of sheets currently held by EthicalArmor GUI."""
from __future__ import annotations
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
from math import isfinite
from pathlib import Path
import json
import os
import re
import tempfile
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

BAD_NAME = re.compile(r"[\\/*?:\[\]\x00-\x1f]")
MAX_ROWS = 1_048_576
MAX_COLS = 16_384

def _clean(value):
    if value is None or value is pd.NA:
        return None
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    if isinstance(value, (datetime, date)):
        if getattr(value, 'tzinfo', None) is not None:
            value = value.isoformat()
        else:
            return value
    if isinstance(value, Decimal):
        value = str(value)
    if isinstance(value, float) and not isfinite(value):
        return None
    try:
        if bool(pd.isna(value)):
            return None
    except (ValueError, TypeError):
        pass
    if isinstance(value, bytes):
        value = value.hex()
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub(' ', value)[:32767]
        if value.startswith(('=', '+', '-', '@', '\t', '\r')):
            value = "'" + value
        return value
    if isinstance(value, (int, float, bool)):
        return value
    try:
        scalar = value.item()
        if scalar is not value:
            return _clean(scalar)
    except (ValueError, TypeError, AttributeError):
        pass
    return _clean(str(value))

def _sheet_name(label, used):
    name = BAD_NAME.sub('_', str(label)).strip("' ")[:31] or 'Sheet'
    out = name
    i = 2
    while out.casefold() in used:
        suffix = '_' + str(i)
        out = name[:31 - len(suffix)] + suffix
        i += 1
    used.add(out.casefold())
    return out

def export_sheets_xlsx(file_path, sheets):
    if not sheets:
        raise ValueError('No sheets to export')
    dest = Path(file_path).expanduser()
    if dest.suffix.lower() != '.xlsx':
        dest = dest.with_suffix('.xlsx')
    book = Workbook(write_only=True)
    used = {'export_manifest'}
    manifests = []
    tmp_path = None
    total = 0
    try:
        for original, frame in sheets.items():
            if not isinstance(frame, pd.DataFrame) or frame.shape[1] == 0:
                continue
            if frame.shape[1] > MAX_COLS:
                raise ValueError(f'Too many columns: {original}')
            header = [_clean(c) for c in frame.columns]
            ids = [bool(re.search(r'imei|imsi|msisdn|iccid|subscriber|phone|mobile|a.party|b.party|cgi|cell.id', str(c), re.I)) for c in frame.columns]
            ws = None
            part = 1
            count = 0
            for row in frame.itertuples(index=False, name=None):
                if ws is None or count >= MAX_ROWS - 1:
                    name = _sheet_name(f'{original} ({part})' if part > 1 else original, used)
                    ws = book.create_sheet(name)
                    ws.append(header)
                    manifests.append([_clean(str(original)), _clean(name), 0])
                    count = 0
                    part += 1
                values = [_clean(v) for v in row]
                ws.append([str(v) if ids[j] and v is not None else v for j, v in enumerate(values)])
                count += 1
                total += 1
                manifests[-1][2] += 1
            if ws is None:
                name = _sheet_name(original, used)
                ws = book.create_sheet(name)
                ws.append(header)
                manifests.append([_clean(str(original)), _clean(name), 0])
        if not manifests:
            raise ValueError('No valid sheets to export')
        meta = book.create_sheet('Export_Manifest')
        meta.append(['GUI Tab', 'Excel Sheet', 'Rows'])
        for item in manifests:
            meta.append(item)
        meta.append(['Source scope', 'GUI-loaded data; preview tabs may be truncated', ''])
        meta.append(['Export time', datetime.now().astimezone().isoformat(), ''])
        with tempfile.NamedTemporaryFile(dir=str(dest.parent), suffix='.xlsx', prefix='.ethicalarmor-', delete=False) as f:
            tmp_path = Path(f.name)
        book.save(str(tmp_path))
        os.replace(tmp_path, dest)
        tmp_path = None
        h = sha256()
        with dest.open('rb') as f:
            for chunk in iter(lambda: f.read(4*1024*1024), b''):
                h.update(chunk)
        return {'path':str(dest), 'sha256':h.hexdigest(), 'sheets':len(manifests), 'rows':total}
    finally:
        book.close()
        if tmp_path is not None:
            tmp_path.unlink(missing_ok=True)
